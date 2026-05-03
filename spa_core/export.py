from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd

from .models import SpaResult
from .segmentation import EVENT_COLUMNS, TOTAL_COLUMNS


SHEET_SPEECH = "Speech Statistics"
SHEET_PAUSE = "Pause Statistics"
SHEET_TOTAL = "Total Statistics"

SEGMENT_AUDIT_COLUMNS = [
    "file_name",
    "source_path",
    "segment_number",
    "segment_type",
    "pause_position",
    "onset_seconds_absolute",
    "offset_seconds_absolute",
    "duration_seconds",
    "audit_json",
]


def _with_filename(frame: pd.DataFrame, filename: str, columns: list[str]) -> pd.DataFrame:
    if frame.empty:
        return pd.DataFrame(columns=["File Name", *columns])
    out = frame.copy()
    out.insert(0, "File Name", filename)
    return out[["File Name", *columns]]


def _write_sheet(workbook, sheet_name: str, title: str, frame: pd.DataFrame, append: bool) -> None:
    if sheet_name in workbook.sheetnames and not append:
        del workbook[sheet_name]
    if sheet_name not in workbook.sheetnames:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append([title])
        sheet.append(list(frame.columns))
    else:
        sheet = workbook[sheet_name]
        if sheet.max_row == 1 and not any(cell.value for cell in sheet[1]):
            sheet.delete_rows(1)
            sheet.append([title])
            sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        sheet.append(list(row))


def export_excel(result: SpaResult, output_path: str | Path, append: bool = False) -> Path:
    """Write MATLAB-style Total/Speech/Pause Statistics sheets to .xlsx."""

    try:
        from openpyxl import Workbook, load_workbook
    except Exception as exc:  # pragma: no cover - optional dependency import path
        raise RuntimeError("openpyxl is required to export Excel files. Install requirements.txt first.") from exc

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    filename = result.signal_path.name

    total = _with_filename(result.total_matrix, filename, TOTAL_COLUMNS)
    speech = _with_filename(result.speech_matrix, filename, EVENT_COLUMNS)
    pause = _with_filename(result.pause_matrix, filename, EVENT_COLUMNS)

    if append and output_path.exists():
        workbook = load_workbook(output_path)
    else:
        workbook = Workbook()
        default = workbook.active
        workbook.remove(default)

    _write_sheet(workbook, SHEET_TOTAL, "TOTAL STATISTICS", total, append=append)
    _write_sheet(workbook, SHEET_SPEECH, "SPEECH STATISTICS", speech, append=append)
    _write_sheet(workbook, SHEET_PAUSE, "PAUSE STATISTICS", pause, append=append)
    workbook.save(output_path)
    return output_path


def segment_audit_frame(result: SpaResult, audit_by_segment_number: dict[int, dict] | None = None) -> pd.DataFrame:
    """Build a compact chronological speech/pause segment table for RA audit."""

    audit_by_segment_number = audit_by_segment_number or {}
    analysis_start = int(result.analysis_region[0])
    sample_rate = float(result.sample_rate)
    rows: list[dict[str, object]] = []

    def add_rows(events: np.ndarray, segment_type: str) -> None:
        for start, end in np.asarray(events, dtype=int):
            absolute_start = analysis_start + int(start)
            absolute_end = analysis_start + int(end)
            if absolute_end <= absolute_start:
                continue
            rows.append(
                {
                    "file_name": result.signal_path.name,
                    "source_path": str(result.signal_path),
                    "_absolute_start_sample": absolute_start,
                    "segment_type": segment_type,
                    "pause_position": 0,
                    "onset_seconds_absolute": absolute_start / sample_rate,
                    "offset_seconds_absolute": absolute_end / sample_rate,
                    "duration_seconds": (absolute_end - absolute_start) / sample_rate,
                    "audit_json": "",
                }
            )

    add_rows(result.speech_events_samples, "speech")
    add_rows(result.pause_events_samples, "pause")
    rows.sort(key=lambda row: (float(row["_absolute_start_sample"]), 0 if row["segment_type"] == "speech" else 1))
    for segment_number, row in enumerate(rows, start=1):
        row["segment_number"] = segment_number
        row["audit_json"] = json.dumps(
            audit_by_segment_number.get(segment_number, {"selected_effects": []}),
            sort_keys=True,
        )
        del row["_absolute_start_sample"]
    if rows and rows[0]["segment_type"] == "pause":
        rows[0]["pause_position"] = "leading"
    if rows and rows[-1]["segment_type"] == "pause":
        rows[-1]["pause_position"] = "leading_trailing" if len(rows) == 1 else "trailing"
    return pd.DataFrame(rows, columns=SEGMENT_AUDIT_COLUMNS)


def export_segment_audit_csv(
    result: SpaResult,
    output_path: str | Path,
    audit_by_segment_number: dict[int, dict] | None = None,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    frame = segment_audit_frame(result, audit_by_segment_number)
    frame.to_csv(output_path, index=False)
    return output_path
