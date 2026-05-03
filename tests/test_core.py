from __future__ import annotations

import importlib.util
import json

import numpy as np
from scipy.io import wavfile

from spa_core.audio import read_wav
from spa_core.export import SEGMENT_AUDIT_COLUMNS, export_excel, segment_audit_frame
from spa_core.models import SpaSettings
from spa_core.processing import amplitude_envelope
from spa_core.segmentation import detect_events, run_spa
from spa_core.thresholds import adaptive_threshold_curve, automatic_static_threshold


def _synthetic_wav(tmp_path):
    fs = 2000
    t = np.arange(0, 2.0, 1 / fs)
    audio = np.zeros_like(t)
    audio += 0.003 * np.random.default_rng(42).normal(size=t.size)
    speech_a = (t >= 0.25) & (t < 0.75)
    speech_b = (t >= 1.10) & (t < 1.65)
    audio[speech_a] += 0.35 * np.sin(2 * np.pi * 180 * t[speech_a])
    audio[speech_b] += 0.30 * np.sin(2 * np.pi * 210 * t[speech_b])
    path = tmp_path / "synthetic.wav"
    wavfile.write(path, fs, (audio * 32767).astype(np.int16))
    return path, fs


def test_amplitude_envelope_normalizes_to_0_100():
    fs = 2000
    t = np.arange(0, 1.0, 1 / fs)
    audio = 0.25 * np.sin(2 * np.pi * 120 * t)
    _, _, _, normalized = amplitude_envelope(audio, fs)
    assert normalized.shape == audio.shape
    assert np.nanmin(normalized) >= 0
    assert np.nanmax(normalized) <= 100.000001
    assert np.isclose(np.nanmax(normalized), 100.0)


def test_static_threshold_formula_is_noise_mean_plus_sd_multiplier():
    envelope = np.array([1, 2, 3, 4, 100, 100], dtype=float)
    threshold, mean_value, std_value = automatic_static_threshold(envelope, (0, 4), 3)
    assert np.isclose(mean_value, 2.5)
    assert np.isclose(std_value, np.std(envelope[:4], ddof=1))
    assert np.isclose(threshold, mean_value + std_value * 3)


def test_adaptive_threshold_uses_35_blocks_and_block_std_term():
    values = np.linspace(0, 10, 350)
    curve = adaptive_threshold_curve(values, noise_mean=2.0, noise_std=1.0, sd_multiplier=3.0)
    assert curve.shape == values.shape
    assert np.all(np.isfinite(curve))
    assert np.nanmean(curve) > 5.0


def test_adaptive_threshold_drops_matlab_remainder_samples():
    values = np.linspace(0, 10, 361)
    curve = adaptive_threshold_curve(values, noise_mean=2.0, noise_std=1.0, sd_multiplier=3.0)
    assert curve.shape == (350,)


def test_detect_events_uses_matlab_boundary_convention():
    envelope = np.array([0.5, 2, 2, 0.5, 0.5, 2, 0.5, 0.5, 0.5, 2, 2, 0.5], dtype=float)
    threshold = np.ones_like(envelope)
    speech, pause = detect_events(
        envelope,
        threshold,
        sample_rate=1000,
        speech_threshold_ms=3,
        pause_threshold_ms=0,
    )
    np.testing.assert_array_equal(speech, np.array([[1, 2], [9, 11]]))
    np.testing.assert_array_equal(pause, np.array([[3, 8]]))


def test_run_spa_detects_synthetic_speech_and_pause(tmp_path):
    path, fs = _synthetic_wav(tmp_path)
    signal = read_wav(path)
    settings = SpaSettings(
        speech_threshold_ms=25,
        pause_threshold_ms=100,
        sd_multiplier=3,
        threshold_mode="automatic",
        adaptive=False,
        use_full_file=True,
    )
    result = run_spa(
        signal,
        settings=settings,
        noise_region=(0, int(0.15 * fs)),
        analysis_region=(0, len(signal.raw_audio)),
    )
    assert len(result.speech_matrix) >= 2
    assert len(result.pause_matrix) >= 1
    assert result.total_matrix.iloc[0]["Speech_events"] == len(result.speech_matrix)


def test_export_excel_writes_matlab_style_sheets(tmp_path):
    if importlib.util.find_spec("openpyxl") is None:
        return
    path, fs = _synthetic_wav(tmp_path)
    signal = read_wav(path)
    result = run_spa(
        signal,
        settings=SpaSettings(adaptive=False, use_full_file=True),
        noise_region=(0, int(0.15 * fs)),
        analysis_region=(0, len(signal.raw_audio)),
    )
    output_path = export_excel(result, tmp_path / "out.xlsx")
    from openpyxl import load_workbook

    workbook = load_workbook(output_path)
    assert {"Total Statistics", "Speech Statistics", "Pause Statistics"}.issubset(workbook.sheetnames)
    assert workbook["Total Statistics"]["A1"].value == "TOTAL STATISTICS"
    assert workbook["Total Statistics"]["A2"].value == "File Name"


def test_segment_audit_frame_uses_absolute_chronological_rows(tmp_path):
    path, fs = _synthetic_wav(tmp_path)
    signal = read_wav(path)
    result = run_spa(
        signal,
        settings=SpaSettings(adaptive=False, use_full_file=False),
        noise_region=(0, int(0.15 * fs)),
        analysis_region=(int(0.2 * fs), int(1.8 * fs)),
    )
    frame = segment_audit_frame(
        result,
        {
            1: {
                "selected_effects": [{"gui_name": "Environmental noise", "effect": "Traffic"}],
            },
            2: {"selected_effects": []},
        },
    )
    assert list(frame.columns) == SEGMENT_AUDIT_COLUMNS
    assert frame["onset_seconds_absolute"].is_monotonic_increasing
    assert frame["onset_seconds_absolute"].min() >= 0.2
    assert set(frame["segment_type"]).issubset({"speech", "pause"})
    assert (frame[frame["segment_type"] == "speech"]["pause_position"] == 0).all()
    if frame.iloc[0]["segment_type"] == "pause":
        assert frame.iloc[0]["pause_position"] == "leading"
    if frame.iloc[-1]["segment_type"] == "pause":
        assert frame.iloc[-1]["pause_position"] == "trailing"
    assert set(frame.iloc[1:-1]["pause_position"]).issubset({0})

    speech_audit = frame[frame["segment_type"] == "speech"].iloc[0]["audit_json"]
    assert json.loads(frame.iloc[0]["audit_json"])["selected_effects"] == [{"gui_name": "Environmental noise", "effect": "Traffic"}]
    assert speech_audit != "NA"
    assert (frame[frame["segment_type"] == "pause"]["audit_json"] != "NA").all()
